"""Generate fictional FR/EN SOP fixtures and an independent alignment manifest."""
from pathlib import Path
import json
import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from translation_aligner.intake import extract_body

OUT = ROOT / "samples"

# Each item defines a gold correspondence, independent of translation output.
# Multiple paragraphs on either side exercise grouping; sentence boundaries are
# ultimately those returned by Systran, not this document generator.
CASES = [
    ("title", "heading", ["Revue du dossier de fabrication"], ["Manufacturing record review"]),
    ("purpose_heading", "heading", ["Objet"], ["Purpose"]),
    ("purpose", "semantic", ["La présente procédure décrit la revue documentaire du dossier de fabrication avant sa transmission au service qualité."], ["This procedure explains how manufacturing records are checked before they are passed to the quality department."]),
    ("scope_heading", "heading", ["Champ d’application"], ["Scope"]),
    ("scope", "fuzzy", ["Cette procédure s’applique aux dossiers du projet Orionex sur le site Belrive."], ["This procedure applies to records for the Orionex project at the Belrive site."]),
    ("checks_heading", "heading", ["Contrôles préliminaires"], ["Initial checks"]),
    ("exact_batch", "exact_candidate", ["Vérifier le numéro de lot."], ["Check the batch number."]),
    ("exact_date", "exact_candidate", ["Vérifier la date."], ["Check the date."]),
    ("fuzzy_signature", "fuzzy", ["Vérifier que toutes les pages du dossier sont datées et signées par l’opérateur."], ["Check that every page of the record has been dated and signed by the operator."]),
    ("fuzzy_attachment", "fuzzy", ["Joindre une copie du certificat d’analyse au dossier de fabrication."], ["Attach a copy of the certificate of analysis to the manufacturing record."]),
    ("semantic_trace", "semantic", ["Toute correction doit permettre de retrouver l’information initiale et d’identifier son auteur."], ["The original entry must remain traceable whenever a change is made, and the person responsible for that change must be identifiable."]),
    ("repeat_first", "repetition", ["Enregistrer le résultat."], ["Record the result."]),
    ("split", "one_to_many", ["Vérifier la présence des annexes et consigner toute pièce manquante dans la fiche de suivi."], ["Check that the attachments are present.", "List any missing documents on the tracking form."]),
    ("source_only", "omission", ["Pour les dossiers historiques, utiliser la chemise grise conservée dans l’armoire documentaire."], []),
    ("page_two", "page_break", [], []),
    ("review_heading", "heading", ["Traitement des écarts"], ["Handling discrepancies"]),
    ("semantic_hold", "semantic", ["En cas de pièce manquante, suspendre la transmission du dossier jusqu’à sa réception."], ["Do not forward the record while any required document is still outstanding."]),
    ("merge", "many_to_one", ["Identifier l’écart dans la fiche de suivi.", "Informer ensuite le responsable qualité."], ["Log the discrepancy on the tracking form, then notify the quality manager."]),
    ("repeat_second", "repetition", ["Enregistrer le résultat."], ["Record the result."]),
    ("target_only", "addition", [], ["For electronic records, include the repository reference in the handover message."]),
    ("numbers", "numeric_equivalence", ["Le dossier de démonstration comporte 12 annexes et porte la référence BR-204."], ["The demonstration record contains 12 attachments and is identified by reference BR-204."]),
    ("numeric_trap", "reject_numeric_mismatch", ["Le registre de démonstration contient 15 entrées."], ["The demonstration log contains 50 entries."]),
    ("negation_trap", "reject_negation_mismatch", ["Ne pas archiver le dossier incomplet."], ["Archive the incomplete record."]),
    ("conflict_a", "translation_conflict", ["Vérifier le dossier."], ["Check the record."]),
    ("context", "fuzzy", ["Après la revue des annexes, effectuer une dernière vérification documentaire."], ["After reviewing the attachments, carry out a final document check."]),
    ("conflict_b", "translation_conflict", ["Vérifier le dossier."], ["Review the file."]),
    ("terms_heading", "heading", ["Terminologie"], ["Terminology"]),
]
TERMS = [
    ("dossier de fabrication", "manufacturing record"),
    ("certificat d’analyse", "certificate of analysis"),
    ("numéro de lot", "batch number"),
    ("traçabilité", "traceability"),
    ("assurance qualité", "quality assurance"),
]


def make_document(lang):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    for name in ("Title", "Heading 1"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(12)
    section.header.paragraphs[0].text = (
        "SOP DEMO 001 • DOCUMENT FICTIF • EXCLURE EN-TÊTE" if lang == "fr" else
        "SOP DEMO 001 • FICTIONAL DOCUMENT • EXCLUDE HEADER")
    section.header.paragraphs[0].style = "Caption"
    section.footer.paragraphs[0].text = (
        "Usage de test uniquement • EXCLURE PIED DE PAGE" if lang == "fr" else
        "Testing only • EXCLUDE FOOTER")
    section.footer.paragraphs[0].style = "Caption"
    doc.core_properties.title = "SOP DEMO 001"
    doc.core_properties.subject = "Synthetic alignment test fixture; not an operational SOP"
    doc.core_properties.author = "Translation aligner test project"
    return doc


def main():
    OUT.mkdir(exist_ok=True)
    docs = {lang: make_document(lang) for lang in ("fr", "en")}
    manifest = []
    for case_id, category, fr, en in CASES:
        if category == "page_break":
            for doc in docs.values():
                doc.add_page_break()
            continue
        for lang, texts in (("fr", fr), ("en", en)):
            for text in texts:
                style = "Title" if case_id == "title" else "Heading 1" if category == "heading" else "Normal"
                docs[lang].add_paragraph(text, style)
        manifest.append({"case": case_id, "category": category, "fr": fr, "en": en})
    for lang, column in (("fr", 0), ("en", 1)):
        table = docs[lang].add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "Code" if lang == "fr" else "Code"
        table.rows[0].cells[1].text = "Terme" if lang == "fr" else "Term"
        for index, terms in enumerate(TERMS, 1):
            cells = table.add_row().cells
            cells[0].text, cells[1].text = f"T{index:02d}", terms[column]
        for row in table.rows:
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement("w:cantSplit"))
        docs[lang].save(OUT / ("SOP_DEMO_001.docx" if lang == "fr" else "SOP_DEMO_001-EN.docx"))
    # Locate gold text in the actual output, using occurrence queues for repeats.
    for lang, name in (("fr", "SOP_DEMO_001.docx"), ("en", "SOP_DEMO_001-EN.docx")):
        locations = {}
        blocks = extract_body((OUT / name).read_bytes())
        for block in blocks:
            locations.setdefault(block.text, []).append(block.position)
        for case in manifest:
            case[lang + "_positions"] = [locations[text].pop(0) for text in case[lang]]
        assert not any("EXCLURE" in b.text or "EXCLUDE" in b.text for b in blocks)
    (OUT / "expected_alignment.json").write_text(json.dumps({
        "synthetic": True,
        "stage_labels_are_intentions_not_live_MT_guarantees": True,
        "cases": manifest,
        "dictionary_pairs": [{"fr": fr, "en": en} for fr, en in TERMS],
        "normalization_pairs": [],
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Generated two DOCX files and alignment manifest in {OUT}")


if __name__ == "__main__":
    main()
